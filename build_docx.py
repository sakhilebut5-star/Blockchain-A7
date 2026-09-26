"""
build_docx.py
Builds report/Assignment7_PartA_Report.docx and report/Assignment7_PartsB_C_Graduate.docx.
Formulas are written in LaTeX and converted by pandoc into native (editable) Word equations.
Run after consensus.py all and analysis.py. Requires pandoc and python-docx.
"""
import csv
import json
import os
import subprocess

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

import consensus as cs

HERE = os.path.dirname(os.path.abspath(__file__))
RES, FIG, OUT = (os.path.join(HERE, d) for d in ("results", "figures", "report"))
os.makedirs(OUT, exist_ok=True)
FONT = "Times New Roman"


def load(n):
    with open(os.path.join(RES, n)) as f:
        return json.load(f)


def rows(n):
    with open(os.path.join(RES, n)) as f:
        return list(csv.DictReader(f))


def md_table(header, body, widths=None):
    widths = widths or [1] * len(header)
    out = ["| " + " | ".join(header) + " |", "|" + "|".join("-" * (4 * w) for w in widths) + "|"]
    out += ["| " + " | ".join(str(c) for c in r) + " |" for r in body]
    return "\n".join(out) + "\n"


def fig(name, caption, width_cm):
    return f"![{caption}]({os.path.join(FIG, name)}){{width={width_cm}cm}}\n"


# --------------------------------------------------------------------------- #
# Part A
# --------------------------------------------------------------------------- #
def part_a_md():
    fork, tie, heavy, ds = (load(x) for x in ("fork.json", "tie.json", "heaviest.json", "doublespend.json"))
    depth = {(float(r["q"]), int(r["z"])): r for r in rows("depth_table.csv")}
    fork_rows = [[f"{s['tag']} (t = {s['t']:.0f} s)"] +
                 [f"{s['nodes'][n]['tip']} (h = {s['nodes'][n]['height']})" for n in "ABC"] +
                 ["yes" if s["converged"] else "no"] for s in fork["snapshots"]]
    notes = {"first_seen": "What Bitcoin Core does. Each node is stable, but the network stays divided until the "
                           "next block, and one side then has a one-block reorganisation.",
             "lowest_hash": "Everyone agrees at once. The catch is that a miner who withholds blocks can publish "
                            "only when its hash will win, which helps selfish mining.",
             "random": "Proposed by Eyal and Sirer (2014) to blunt selfish mining. In our run A and B switched to "
                       "C2 while C switched to A2, so the split survived."}
    tie_rows = []
    for pol in ("first_seen", "lowest_hash", "random"):
        d, a = tie[pol]["during_tie"]["nodes"], tie[pol]["after"]
        tie_rows.append([pol, " / ".join(d[n]["tip"] for n in "ABC"),
                         f"all on {a['nodes']['A']['tip']}" if a["converged"] else "split", notes[pol]])
    need = {q: cs.required_depth(q, 1e-3) for q in (0.1, 0.2, 0.3)}
    dep_rows = [[f"{q:.1f}"] + [f"{float(depth[(q, z)]['rosenfeld_premine1']):.4f}" for z in (1, 2, 3, 6)]
                + [str(need[q])] for q in (0.1, 0.2, 0.3)]
    L, W = heavy["length"]["nodes"]["B"], heavy["work"]["nodes"]["B"]

    return f"""---
title: "Assignment 7: Consensus and Forks Lab"
subtitle: "BLCH9X2 Blockchain | Master of Financial Engineering, University of Johannesburg | Group submission, Part A (60%)"
---

# 1. Objective and design

We built a small Nakamoto-consensus simulator with three nodes, A, B and C, running in one process. Each node keeps its own copy of the chain, and those copies are allowed to drift apart. Instead of a single list, every node stores all the valid blocks it has seen as a *block tree*, and a fork-choice rule picks the tip it builds on. The tree is what makes switching possible: a node holding only a list would have nowhere to go when a longer competing branch turned up.

The main experiments use the **longest-chain rule**, as the brief allows. We also implemented heaviest cumulative work, but only as a short extension in Section 3.3. All of the code runs from `consensus_lab.ipynb`. The same code is also saved as `consensus.py`, the file the brief names.

{md_table(["Component", "Responsibility"], [
    ["Tx, Block, mine()", "Simple UTXO-style transactions, a Merkle root and a double-SHA-256 header. Proof of work is real but cheap: 12 leading zero bits, or roughly 4 096 hashes a block."],
    ["apply_block()", "Checks each block against the ledger. Inputs must exist and be unspent, outputs cannot exceed inputs, and the coinbase cannot pay more than subsidy plus fees. A branch is accepted only if its whole history passes."],
    ["Node", "Holds the block tree, an orphan pool for blocks that arrive before their parent, and a mempool. Fork choice is by length (default) or by work (extension), with first-seen, lowest-hash or random tie-breaking. In a reorganisation, transactions from dropped blocks go back to the mempool, or are flagged as reversed if they now conflict."],
    ["Network", "A discrete-event gossip model. Each link has its own latency and, optionally, bandwidth, so big blocks travel slowly. We can also split the network and hold messages until it heals. A node passes every block it accepts on to its other peers."]], [4, 16])}

**Proof of work and fork choice.** A header hash $H$ is valid when it is below the target set by the difficulty $b$ (the number of leading zero bits). The work in a block is the number of hashes you would expect to try before finding it:

$$H < 2^{{256-b}}, \\qquad w = 2^{{b}}.$$

For a leaf $x$ of the tree, write $h(x)$ for its height and $W(x) = \\sum_{{i \\in \\text{{path}}(x)}} 2^{{b_i}}$ for the total work along its path. The two rules then choose

$$\\text{{tip}}_{{\\text{{length}}}} = \\arg\\max_{{x}} \\; h(x), \\qquad \\text{{tip}}_{{\\text{{work}}}} = \\arg\\max_{{x}} \\; W(x),$$

Ties go to whichever tie-break policy is set. If every block has the same $b$, then $W(x) = (h(x)+1)\\,2^{{b}}$, so both rules pick the same tip.

Running the notebook from top to bottom replays every scenario, shows the event logs and figures, and saves the results as JSON in `results/`. Eight tests, one for each claim we make below, run with `pytest -q`. Time is simulated and measured in seconds. We chose delays that keep the logs easy to follow; they are not meant to copy Bitcoin's ten-minute rhythm.

# 2. Creating and resolving a fork

We start with a shared chain G, A1, B2. At t = 19 s, A and C both mine a valid block on top of B2, so there are now two blocks at height 3. The links are deliberately uneven (A-B takes 1 s, B-C 3 s and A-C 6 s), which means B hears about A3 before C3.

With first-seen tie-breaking, a node sticks with the first block it received when an equally long rival arrives. The network ends up in two camps, A and B on A3 and C on C3, and stays that way until someone mines again. B does so at t = 29 s, building B4 on A3. When B4 reaches C at t = 32 s, the A3 branch is one block longer (height 4 against 3). C drops C3, adopts A3 and B4, and all three nodes agree again. C3 is now stale, and the coinbase C paid itself in it no longer exists on the ledger.

{md_table(["Stage (time)", "A tip", "B tip", "C tip", "Converged"], fork_rows, [8, 3, 3, 3, 3])}

{fig("fig1_block_tree.png", "Figure 1. Block tree held by every node after resolution. The shaded block C3 is valid but off the best chain.", 13.5)}

In the event log this appears as `t=22 B: TIE C3 vs tip A3; keep A3` and, ten seconds later, `t=32 C: REORG depth 1 at fork point B2: drop [C3] adopt [A3, B4]`. All blocks here share one difficulty, so the heaviest-work rule would have made exactly the same choice.

# 3. Edge cases

## 3.1 Ties

Two tips of equal length make a tie. Nothing has gone wrong at that point. Both branches are valid, and the next block will decide between them. We re-ran one fork (A2 and C2, both on B1) under each tie-break policy.

{md_table(["Tie-break", "Tips during tie (A / B / C)", "After next block", "Comment"], tie_rows, [3, 4, 3, 12])}

All three policies agreed after one more block. For a tie to last, the two branches would have to keep growing at exactly the same pace, and that gets less likely with every block. The flip side is that one confirmation is weak evidence: the block that confirms a payment may be one half of a tie the merchant has no way of seeing.

## 3.2 Delayed and out-of-order block arrival

**Child before parent.** Here the nodes are connected in a line (A-B-C) over 50 kB/s links. A block of size $s$ sent over a link with latency $\\ell$ and bandwidth $\\beta$ takes $\\delta = \\ell + s/\\beta$ to arrive. A mines a 400 kB block, A2big, which needs about 8 s per hop. B then mines a small block on top of it, B3small, and the small block reaches C first, at t = 10.1 s. C cannot check B3small without the parent's UTXO state, so it keeps it in the orphan pool instead of throwing it away. When A2big arrives at t = 18.0 s, C connects the parent and then the child straight away. Decker and Wattenhofer (2013) point to exactly this effect, large blocks travelling slowly, as the main reason forks happen.

**Partition and late delivery.** Next we cut the network into {{A, B}} and {{C}}. The larger side mines three blocks (A2, B3, A4) and C mines two (C2, C3), with every message across the cut held back. Once the partition heals, C sees a longer branch and reorganises two blocks deep, dropping C2 and C3. A and B receive C's blocks late, file them as a side branch and stay where they are. Late news can move a node onto a longer chain, but never back onto a shorter one.

## 3.3 Extension: longest chain versus heaviest work

Branch A has three easy blocks ($b = 8$, total work $3 \\times 2^{{8}} = 768$). Branch C has two hard ones ($b = 13$, total work $2 \\times 2^{{13}} = 16\\,384$). Under the length rule, every node settles on {L['tip']} at height {L['height']}. Under the work rule, they settle on {W['tip']} at height {W['height']}, with cumulative work of {W['cum_work']:,}. With a fixed difficulty, as in all our other scenarios, the two rules cannot disagree. They only part ways when blocks differ in difficulty, and then counting blocks lets an attacker win with a pile of cheap ones. That is why Bitcoin compares total chainwork. In the lab, miners pick their own difficulty to make this visible; in Bitcoin, retargeting sets it.

# 4. Confirmation depth for a R10 000 payment

**Definition.** Take a payment in block $x$ and a node whose tip is at height $h_{{\\text{{tip}}}}$. If $x$ is on that node's best chain, the payment has

$$z = h_{{\\text{{tip}}}} - h(x) + 1$$

confirmations; otherwise it has none. The count depends on who you ask. During the fork in Section 2, a transaction that only appeared in C3 had one confirmation at C and zero at A and B.

**Demonstration.** C, playing the attacker, pays merchant B R10 000 from a genesis coin, and honest miner A puts the payment into block A1pay. Meanwhile C mines in private from the parent block, with a transaction that sends the same coin back to itself. The merchant ships at $z = {ds['z']}$. C then releases its private branch, which is one block longer. B reorganises (depth {ds['reorgs'][0]['depth']}), the payment is marked as reversed, and the merchant's balance falls from R10 000 to R{ds['merchant_balance_after']}. We scripted C's luck here. A real attacker would have to out-mine everyone else for as long as the race lasts.

**Probability.** Write $q$ for the attacker's share of hash power and $p = 1 - q$. Nakamoto (2008) approximates how far the attacker gets, while the merchant waits for $z$ blocks, with a Poisson distribution of mean $\\lambda = zq/p$:

$$P_{{N}}(q,z) = 1 - \\sum_{{k=0}}^{{z}} \\frac{{\\lambda^{{k}} e^{{-\\lambda}}}}{{k!}} \\left(1 - \\left(\\frac{{q}}{{p}}\\right)^{{z-k}}\\right).$$

Rosenfeld (2014) works it out exactly with a negative binomial. He also gives the attacker one block mined in advance, which is the cautious assumption:

$$P_{{R}}(q,z) = 1 - \\sum_{{m=0}}^{{z}} \\binom{{m+z-1}}{{m}} \\left(p^{{z}} q^{{m}} - p^{{m}} q^{{z}}\\right), \\qquad q < p.$$

If the attacker has no head start, it has to finish strictly ahead, and the probability becomes

$$P_{{S}}(q,z) = 1 - \\sum_{{m=0}}^{{z}} \\binom{{m+z-1}}{{m}} p^{{z}} q^{{m}} \\left(1 - \\left(\\frac{{q}}{{p}}\\right)^{{z-m+1}}\\right).$$

We coded all three and checked them against a Monte Carlo race; Figure 2 shows how closely they match. Without the head start, the attacker's chances are roughly 3 to 6 times lower (`results/depth_table.csv`). For a tolerated reversal probability $\\varepsilon$, the depth a policy needs is

$$z^{{*}}(q,\\varepsilon) = \\min\\{{\\, z \\ge 1 : P_{{R}}(q,z) \\le \\varepsilon \\,\\}}.$$

{fig("fig2_confirmation_depth.png", "Figure 2. Probability that an attacker with hash share q eventually reverses a payment accepted after z confirmations.", 11.5)}

{md_table(["Attacker share q", "z = 1", "z = 2", "z = 3", "z = 6", "z* for ε = 0.1%"], dep_rows)}

**Why the answer depends on assumptions.** It is tempting to compare R10 000 with the block subsidy and conclude that nobody would bother. That comparison does not settle anything. An attacker might double-spend many payments in one go, might have reasons that have nothing to do with the money, might rent hash power, or might already own mining equipment. The depth really comes down to two numbers the institution has to choose: the attacker share $q$ it plans for and the reversal probability $\\varepsilon$ it can live with. With $\\varepsilon = 0.1\\%$, for example, a 10% attacker means waiting for 6 confirmations and a 20% attacker means 13.

Zero-confirmation payments face a different problem. In a *race attack*, the payer sends one transaction to the merchant and a conflicting one to the miners. No hash power is needed at all.

**What depth means for R10 000.** For a R10 000-equivalent payment, confirmation depth should be chosen according to the assumed attacker capability, observed network conditions, the value of the transaction and how much reversal risk the institution will accept. Accepting at zero confirmations gives the weakest protection, and each additional confirmation lowers the risk of a reorganisation further.

# 5. Limitations and extensions

A few simplifications are worth stating. Transactions are not signed, since we covered signatures in the ECDSA lab. Difficulty never retargets. The gossip model leaves out inventory messages, compact blocks and peer selection. And the double-spend scenario hard-codes the attacker's luck instead of sampling it, which is why the probabilities come from the formulas and the Monte Carlo runs. With more time we would add retargeting, headers-first sync and a selfish-mining agent. The last of these would let us measure how much each tie-break policy helps a miner who withholds blocks.

# 6. Reproducibility and verification

Every claim above is backed by an automated test, so a marker can re-run the whole set in about two seconds.

{md_table(["Test (tests/test_consensus.py)", "Claim verified"], [
    ["test_pow_and_tamper_rejected", "A block whose header was altered after mining is rejected; a valid block extends the tip."],
    ["test_fork_resolves_to_single_tip", "Section 2: split after propagation, single tip after the next block, one depth-1 reorganisation."],
    ["test_first_seen_tie_keeps_split_until_next_block", "Section 3.1: first-seen stays split, lowest-hash converges immediately, all policies converge after one more block."],
    ["test_orphan_pool_connects_late_parent", "Section 3.2: child-before-parent is parked and later connected; partition heals to one tip."],
    ["test_heaviest_work_beats_longest", "Section 3.3: the two rules select different branches."],
    ["test_double_spend_reverses_payment_at_z1", "Section 4: the R10 000 payment is reversed and the merchant balance returns to zero."],
    ["test_in_block_double_spend_invalid", "A block spending the same coin twice is invalid."],
    ["test_formulas_agree_with_monte_carlo", "Rosenfeld and strict formulas match simulation within 1.5 percentage points; Nakamoto matches the whitepaper table (q = 0.1, z = 5)."]], [8, 12])}

To reproduce everything, install the requirements (`pip install -r requirements.txt`), open `consensus_lab.ipynb` and run all cells.

# 7. Contribution statement

Each member read through all of the code and the report. The names and the split below are placeholders to update before submission.

{md_table(["Member", "Contribution"], [
    ["Sakhile", "Node class, fork-choice and reorganisation logic; Sections 1 and 2; final editing."],
    ["[Member 2]", "Network simulator (latency, bandwidth, partitions); delayed-arrival and tie scenarios; Section 3."],
    ["[Member 3]", "Confirmation-depth analytics (Nakamoto, Rosenfeld, Monte Carlo); double-spend scenario; Section 4."],
    ["[Member 4]", "pytest suite, figures and README; limitations and extensions."]], [3, 17])}

# References

Decker, C. and Wattenhofer, R. (2013) 'Information propagation in the Bitcoin network', *13th IEEE International Conference on Peer-to-Peer Computing (P2P 2013)*, Trento.

Eyal, I. and Sirer, E.G. (2014) 'Majority is not enough: Bitcoin mining is vulnerable', *Financial Cryptography 2014*, LNCS 8437, pp. 436-454.

Nakamoto, S. (2008) *Bitcoin: A peer-to-peer electronic cash system*.

Rosenfeld, M. (2014) 'Analysis of hashrate-based double spending', arXiv:1402.2009.
"""


# --------------------------------------------------------------------------- #
# Parts B and C
# --------------------------------------------------------------------------- #
def parts_bc_md():
    pol = {(float(r["T"]), float(r["delay_s"]), float(r["q"])): r
           for r in rows("relay_policy.csv") if r["kind"] == "policy"}
    b6, b40 = pol[(600.0, 5.0, 0.1)], pol[(600.0, 40.0, 0.1)]
    f6, f40 = pol[(60.0, 5.0, 0.25)], pol[(60.0, 40.0, 0.25)]
    bands = [("< R1 000", 1e-2), ("R1 000 to R10 000", 1e-3), ("R10 000 to R100 000", 1e-4), ("> R100 000", 1e-5)]
    band_rows = [[b, f"{e:.3%}"] + [str(cs.required_depth(q, e)) for q in (0.05, 0.10, 0.25)] + ["1 committed block"]
                 for b, e in bands]
    return f"""---
title: "Assignment 7: Graduate Extensions"
subtitle: "Part B (20%): block relay time and confirmation policy for remittances | Part C (20%): group memo on confirmation depth for cross-border retail payments"
---

# Part B. Block relay time versus confirmation policy

In Mastering Bitcoin (Antonopoulos and Harding, 2023, Ch. 11), the blockchain is a tree of blocks linked by header hashes. Each node checks blocks for itself and follows the best branch it knows about, so a block only counts at a node once it has arrived there and passed validation. We wanted to know how the time this relay takes should affect the number of confirmations a remittance operator asks for.

The connection is forks. While a new block is still spreading, other miners are still working on the old tip, and anything they find becomes a competing block. Work spent on the branch that loses is wasted. An attacker mining in private never competes with itself, so in effect it faces a smaller honest network.

**Model.** Treat block arrivals as a Poisson process with mean interval $T$. If a block needs time $d$ to reach the other miners, the chance that a rival appears in that window is about

$$f \\approx 1 - e^{{-d/T}},$$

This is the simplest version of Decker and Wattenhofer's model, with one fixed delay for everyone. Our simulation has $n = 10$ equal miners, and the miner who found a block cannot compete with itself, so we compare the results with $f_n = 1 - e^{{-\\frac{{n-1}}{{n}}\\,d/T}}$. Forks slow honest chain growth by a factor of $(1 - f)$, which pushes the attacker's effective share up to

$$q_{{\\text{{eff}}}} = \\frac{{q}}{{q + (1-q)(1-f)}}.$$

We set the policy depth to $z^{{*}} = \\min\\{{z \\ge 1 : P_{{R}}(q_{{\\text{{eff}}}}, z) \\le \\varepsilon\\}}$, with $\\varepsilon = 0.1\\%$ and $P_R$ from Rosenfeld (2014). The expected wait is then

$$\\mathbb{{E}}[\\text{{wait}}] = \\frac{{z^{{*}}\\, T}}{{1 - f}}.$$

Once $q_{{\\text{{eff}}}}$ reaches $\\tfrac{{1}}{{2}}$, no depth is enough.

{fig("fig3_relay_vs_policy.png", "Figure 3. Relay time drives the fork rate (a), which raises the depth needed for a 0.1% reversal bound (b) and the time a remittance recipient waits (c). Solid: Bitcoin (T = 600 s). Dashed: a hypothetical 60-second chain.", 16.5)}

{fig("fig4_remittance_timeline.png", "Figure 4. Where relay time sits in a Bitcoin remittance: it is a small slice of each confirmation interval, so it changes risk mainly through forks, not through waiting time.", 15)}

**Reading the diagram.** On Bitcoin, $d$ is tiny next to $T$. If relay slows from 5 to 40 seconds, the stale rate goes from {float(b6['stale_sim_or_f']):.1%} to {float(b40['stale_sim_or_f']):.1%}, yet $z^{{*}}$ stays at {b6['z_req']} for $q = 0.1$ and the wait only grows from {float(b6['wait_min']):.0f} to {float(b40['wait_min']):.0f} minutes. A chain with 60-second blocks behaves very differently. The same slowdown lifts its stale rate from {float(f6['stale_sim_or_f']):.0%} to {float(f40['stale_sim_or_f']):.0%}, and for $q = 0.25$ it pushes $z^{{*}}$ from {f6['z_req']} to {f40['z_req']}. Slow relay down further and no depth is safe, because $q_{{\\text{{eff}}}}$ closes in on one half. Shorter block times only speed up settlement while relay stays well under $T$.

**What this means for remittance.** A policy should be stated in minutes as well as blocks, because ten-minute blocks with slow relay are safer than one-minute blocks with slow relay. Operators should watch how many stale blocks their own nodes see, since a rise is an early sign that $z^{{*}}$ needs to go up. Faster relay, through compact block relay (BIP 152) and high-bandwidth peers, also improves security: it keeps $d/T$ small and so keeps $q_{{\\text{{eff}}}}$ close to $q$. Finally, an operator should connect to plenty of well-connected peers, so that its own view of the chain is not the one lagging behind.

# Part C. Memo: confirmation depth for cross-border retail payments

**To:** Payments product committee, [SA remittance provider]

**From:** Group [number], BLCH9X2

**Date:** September 2026

**Subject:** How many confirmations before a R10 000 cross-border payment is final?

**Recommendation.** For a R10 000-equivalent payment, confirmation depth should be chosen according to the assumed attacker capability, observed network conditions, the value of the transaction and how much reversal risk the institution will accept. Accepting at zero confirmations gives the weakest protection, and each additional confirmation lowers the risk of a reorganisation further. In practice, we suggest the committee (i) sets an assumed attacker share and an acceptable reversal probability for each value band, (ii) works out the Bitcoin depth from those two figures, (iii) raises the depth whenever monitoring picks up unusual fork activity, and (iv) on a Byzantine-fault-tolerant (BFT) chain, releases funds after the first committed block, as long as the validator set is known and no more than a third of validators can be faulty.

**1. Why Bitcoin finality is probabilistic.** Decker and Wattenhofer (2013) measured how quickly blocks spread across the Bitcoin network. The median was about 6.5 seconds and the mean 12.6 seconds, with a long tail: reaching 95% of nodes took roughly 40 seconds. They showed that this delay explains the fork rate they saw, about 1.7% of blocks, and that most of it comes from each hop verifying and forwarding large blocks. Because forks happen, a confirmed block can still be replaced, so every depth $z$ leaves some reversal probability $P(q,z) > 0$. Our Part A simulation shows the same thing on a small scale. A slow block led to an orphan and an out-of-order reorganisation, and a lucky private chain reversed a R10 000 payment that had one confirmation. The authors also point out that propagation delay helps double-spenders, because different parts of the network briefly see different transactions.

**2. Why BFT finality is deterministic, with conditions.** Castro and Liskov's (1999) Practical Byzantine Fault Tolerance (PBFT) runs on $n = 3f + 1$ known replicas. A request commits after three phases (pre-prepare, prepare and commit), and each phase needs a quorum of $2f + 1$. Any two quorums share

$$2(2f+1) - (3f+1) = f + 1$$

replicas, so at least one honest replica sits in both. As long as no more than $f$ replicas are faulty, two conflicting blocks cannot both commit. Timing affects liveness but not safety, so a committed block cannot be reorganised and one confirmation is enough. The risk does not disappear, though. It becomes a question of who the validators are: if more than a third of them collude, they can commit conflicting histories. PBFT's message cost also grows with the square of the number of validators, which confines it to permissioned networks. That suits a regulated payment corridor, but not an open network.

{md_table(["Value band (ZAR)", "Tolerated reversal probability (illustrative)", "Bitcoin z* if q = 0.05", "Bitcoin z* if q = 0.10", "Bitcoin z* if q = 0.25", "BFT chain"], band_rows, [4, 4, 3, 3, 3, 4])}

Depths use Rosenfeld (2014) with $T = 10$ minutes. The tolerances are illustrative settings for risk appetite, not recommendations. The BFT column assumes at most $f$ of the $3f + 1$ validators are faulty.

**3. Choosing a depth for R10 000.** The table shows why we cannot name one objectively safe number. Say the committee accepts a 0.1% chance of reversal for this band. Against a 5% attacker, that means 4 confirmations (about 40 minutes). Against 10% it means 6, roughly an hour, and against 25% it means 20. Comparing R10 000 with the block subsidy does not resolve this, because an attacker may double-spend many payments at once, may have motives beyond the money, may rent hash power or may already run mining equipment. Part A also showed that a single block can be one side of a tie the merchant cannot see, so one confirmation is only a small step up from none. What the committee actually has to decide is which $q$ and which $\\varepsilon$ it is willing to defend. The depth follows from those.

**4. Customer and compliance implications.** People receiving money through the corridor experience finality as a wait, so the committee should publish expected release times for each band. On Bitcoin these come to roughly $z \\times 10$ minutes (Part B). The provider should also track stale blocks and peer lag on its own nodes and move to a higher band automatically when conditions get worse. None of this replaces compliance. Crypto assets are financial products under FAIS, the Financial Intelligence Centre's travel-rule directive requires originator and beneficiary details for crypto transfers, and exchange-control reporting to the South African Reserve Bank still applies to the rand leg.

**5. Alternatives.** For frequent small remittances, a payment channel such as Lightning settles in seconds, with finality at the channel level. Another route is a permissioned BFT ledger shared by licensed providers, which gives one-block finality and clear accountability. Either option changes the question. Instead of asking how many blocks to wait for, the committee would be asking whom it trusts and how that trust is enforced.

# References

Antonopoulos, A.M. and Harding, D.A. (2023) *Mastering Bitcoin: Programming the Open Blockchain*, 3rd edn. Sebastopol: O'Reilly, Ch. 11.

Castro, M. and Liskov, B. (1999) 'Practical Byzantine fault tolerance', *Proceedings of the 3rd Symposium on Operating Systems Design and Implementation (OSDI '99)*, New Orleans, pp. 173-186.

Decker, C. and Wattenhofer, R. (2013) 'Information propagation in the Bitcoin network', *13th IEEE International Conference on Peer-to-Peer Computing (P2P 2013)*, Trento.

Nakamoto, S. (2008) *Bitcoin: A peer-to-peer electronic cash system*.

Rosenfeld, M. (2014) 'Analysis of hashrate-based double spending', arXiv:1402.2009.
"""


# --------------------------------------------------------------------------- #
# Styling
# --------------------------------------------------------------------------- #
def make_reference(path):
    subprocess.run(["pandoc", "-o", path, "--print-default-data-file", "reference.docx"], check=True)
    d = Document(path)
    black = RGBColor(0, 0, 0)
    for st in d.styles:
        try:
            f = st.font
        except AttributeError:
            continue
        if f is None:
            continue
        f.name = FONT
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), FONT)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rf.get(qn(a)) is not None:
                del rf.attrib[qn(a)]
        if st.name not in ("Hyperlink",):
            f.color.rgb = black
            c = rpr.find(qn("w:color"))
            if c is not None:
                for a in ("w:themeColor", "w:themeShade", "w:themeTint"):
                    if c.get(qn(a)) is not None:
                        del c.attrib[qn(a)]
    dd = d.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    rf = dd.find(qn("w:rFonts"))
    for a in list(rf.attrib):
        del rf.attrib[a]
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    sizes = {"Normal": 10.5, "Body Text": 10.5, "First Paragraph": 10.5, "Compact": 9, "Title": 16,
             "Subtitle": 10.5, "Heading 1": 12.5, "Heading 2": 11, "Image Caption": 9, "Caption": 9}
    by_name = {x.name.lower(): x for x in d.styles}
    for name, sz in sizes.items():
        st = by_name.get(name.lower())
        if st is not None:
            st.font.size = Pt(sz)
            if name.startswith("Heading") or name == "Title":
                st.font.bold = True
                st.font.italic = False
            if name in ("Image Caption", "Caption", "Subtitle"):
                st.font.italic = True
            if name in ("Body Text", "First Paragraph"):
                st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                st.paragraph_format.space_after = Pt(4)
                st.paragraph_format.space_before = Pt(0)
                st.paragraph_format.line_spacing = 1.0
            if name.startswith("Heading"):
                st.paragraph_format.space_before = Pt(8)
                st.paragraph_format.space_after = Pt(3)
    vc = by_name.get("verbatim char")
    if vc is not None:
        vc.font.name = "Consolas"; vc.font.size = Pt(9)
        rf2 = vc.element.get_or_add_rPr().find(qn("w:rFonts"))
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf2.set(qn(a), "Consolas")
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.0))
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    d.save(path)


def border(tag, sz, color):
    e = OxmlElement(f"w:{tag}")
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), "0"); e.set(qn("w:color"), color)
    return e


def style_tables_and_footer(path, footer_text):
    d = Document(path)
    for t in d.tables:
        tblPr = t._tbl.tblPr
        for old in tblPr.findall(qn("w:tblBorders")) + tblPr.findall(qn("w:shd")):
            tblPr.remove(old)
        b = OxmlElement("w:tblBorders")
        b.append(border("top", 8, "000000"))
        b.append(border("bottom", 8, "000000"))
        b.append(border("insideH", 4, "A6A6A6"))
        tblPr.append(b)
        for tag in ("w:tblLook",):
            for old in tblPr.findall(qn(tag)):
                old.set(qn("w:firstRow"), "0"); old.set(qn("w:val"), "0000")
        tw = tblPr.find(qn("w:tblW"))
        if tw is None:
            tw = OxmlElement("w:tblW"); tblPr.append(tw)
        tw.set(qn("w:type"), "pct"); tw.set(qn("w:w"), "5000")
        for i, row in enumerate(t.rows):
            for cell in row.cells:
                tcPr = cell._tc.get_or_add_tcPr()
                for old in tcPr.findall(qn("w:shd")):
                    tcPr.remove(old)
                if i == 0:
                    bb = OxmlElement("w:tcBorders"); bb.append(border("bottom", 6, "000000")); tcPr.append(bb)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for r in p.runs:
                        r.font.size = Pt(9)
                        if i == 0:
                            r.font.bold = True
    for t in d.tables:            # breathing space after each table
        gap = OxmlElement("w:p")
        ppr = OxmlElement("w:pPr"); sp = OxmlElement("w:spacing")
        sp.set(qn("w:before"), "0"); sp.set(qn("w:after"), "60"); sp.set(qn("w:line"), "120")
        sp.set(qn("w:lineRule"), "exact"); ppr.append(sp); gap.append(ppr)
        t._tbl.addnext(gap)
    for p in d.paragraphs:        # centre figures
        if p.style.name in ("Captioned Figure", "Figure") or any(
                r._element.findall(".//" + qn("w:drawing")) for r in p.runs):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ft = d.sections[0].footer.paragraphs[0]
    ft.text = ""
    r = ft.add_run(footer_text + "    Page ")
    r.font.size = Pt(8); r.font.name = FONT
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        run = ft.add_run(); run.font.size = Pt(8); run.font.name = FONT
        if kind:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); run._r.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; run._r.append(it)
    ft.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d.save(path)


def build(md, name):
    ref = os.path.join(OUT, "_reference.docx")
    if not os.path.exists(ref):
        make_reference(ref)
    md_path = os.path.join(OUT, name + ".md")
    with open(md_path, "w") as f:
        f.write(md)
    out = os.path.join(OUT, name + ".docx")
    subprocess.run(["pandoc", md_path, "-f", "markdown+tex_math_dollars", "-t", "docx",
                    "--reference-doc", ref, "-o", out], check=True)
    style_tables_and_footer(out, "BLCH9X2 Assignment 7: Consensus and Forks Lab")
    os.remove(md_path)
    return out


if __name__ == "__main__":
    a = build(part_a_md(), "Assignment7_PartA_Report")
    b = build(parts_bc_md(), "Assignment7_PartsB_C_Graduate")
    os.remove(os.path.join(OUT, "_reference.docx"))
    print(a, b)
